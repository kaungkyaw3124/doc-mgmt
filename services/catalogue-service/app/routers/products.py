import io
import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from decimal import Decimal, InvalidOperation

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Header
from fastapi.responses import Response
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.db import get_db
from app.core.search_client import index_product, remove_product_from_index
from app.core.storage import upload_file, get_presigned_url, download_file_bytes
from app.core.document_client import get_visible_product_ids
from app import models, schemas

router = APIRouter(prefix="/products", tags=["products"])


def _sub_item_counts_map(db: Session, product_ids: list) -> dict:
    """One query for however many products, instead of one query per product."""
    if not product_ids:
        return {}
    rows = (
        db.query(models.ProductSubItem.parent_product_id, func.count(models.ProductSubItem.id))
        .filter(models.ProductSubItem.parent_product_id.in_(product_ids))
        .group_by(models.ProductSubItem.parent_product_id)
        .all()
    )
    return {str(pid): count for pid, count in rows}


def _attach_sub_item_counts(db: Session, products: list):
    counts = _sub_item_counts_map(db, [p.id for p in products])
    for p in products:
        p.sub_item_count = counts.get(str(p.id), 0)
    return products


def _attach_sub_item_count(db: Session, product):
    product.sub_item_count = (
        db.query(models.ProductSubItem).filter_by(parent_product_id=product.id).count()
    )
    return product


def _extract_zip_entries(file_bytes: bytes):
    try:
        zf = zipfile.ZipFile(io.BytesIO(file_bytes))
    except zipfile.BadZipFile:
        raise HTTPException(status_code=400, detail="that catalogue file isn't a valid zip")
    entries = []
    for entry_name in zf.namelist():
        if entry_name.endswith("/"):
            continue  # a directory entry, not a file
        entries.append((entry_name.rsplit("/", 1)[-1], zf.read(entry_name)))
    return entries


def _extract_rar_entries(file_bytes: bytes):
    """
    Shells out to `unar` (installed in the container image) rather than
    relying on a Python RAR library — there's no free, pure-Python RAR
    decoder, and `unar` is a well-supported, actually-free command-line
    tool that handles it reliably.
    """
    tmp_dir = tempfile.mkdtemp(prefix="rar_import_")
    try:
        rar_path = os.path.join(tmp_dir, "archive.rar")
        with open(rar_path, "wb") as f:
            f.write(file_bytes)

        extract_dir = os.path.join(tmp_dir, "extracted")
        os.makedirs(extract_dir, exist_ok=True)

        result = subprocess.run(
            ["unar", "-quiet", "-no-directory-confirmation", "-force-overwrite", "-output-directory", extract_dir, rar_path],
            capture_output=True,
            timeout=60,
        )
        if result.returncode != 0:
            raise HTTPException(
                status_code=400,
                detail="couldn't read that rar file — it may be corrupted, password-protected, or in an unsupported format",
            )

        entries = []
        for root, _dirs, files in os.walk(extract_dir):
            for fname in files:
                with open(os.path.join(root, fname), "rb") as f:
                    entries.append((fname, f.read()))
        return entries
    except subprocess.TimeoutExpired:
        raise HTTPException(status_code=400, detail="that rar file took too long to extract")
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _extract_archive_entries(filename: str | None, file_bytes: bytes):
    """Returns [(basename, bytes), ...] from either a .zip or .rar archive,
    picked by the uploaded file's extension. Falls back to zip if the
    extension is missing or unrecognized — matches this endpoint's
    original behavior before .rar support was added."""
    ext = filename.rsplit(".", 1)[-1].lower() if filename and "." in filename else ""
    if ext == "rar":
        return _extract_rar_entries(file_bytes)
    return _extract_zip_entries(file_bytes)


def _generate_sku(db: Session, category_name: str | None) -> str:
    """
    e.g. COM-0001 for a product in the "Computer" category — using that
    category's manually-set short_term as the prefix (not derived from
    the category name itself). Falls back to "PRD" if no category was
    given, or the category has no short_term set. Finds the next free
    number for that prefix, so it stays sequential per-category.
    """
    prefix = "PRD"
    if category_name:
        category = db.query(models.Category).filter_by(name=category_name).first()
        if category and category.short_term:
            prefix = category.short_term.upper()

    existing_skus = [
        p.sku for p in db.query(models.Product.sku).filter(models.Product.sku.like(f"{prefix}-%")).all()
    ]
    max_n = 0
    for sku in existing_skus:
        match = re.match(rf"^{re.escape(prefix)}-(\d+)$", sku)
        if match:
            max_n = max(max_n, int(match.group(1)))
    return f"{prefix}-{max_n + 1:04d}"


@router.post("", response_model=schemas.ProductOut, status_code=201)
def create_product(payload: schemas.ProductCreate, db: Session = Depends(get_db)):
    sku = payload.sku or _generate_sku(db, payload.category)

    existing = db.query(models.Product).filter_by(sku=sku).first()
    if existing:
        raise HTTPException(status_code=409, detail="sku already exists")

    data = payload.model_dump()
    data["sku"] = sku
    product = models.Product(**data)
    db.add(product)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="sku already exists")
    db.refresh(product)
    index_product(product)
    return _attach_sub_item_count(db, product)


@router.get("", response_model=list[schemas.ProductOut])
def list_products(
    category: str | None = None,
    q: str | None = None,   # simple name search, e.g. ?q=widget
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    query = db.query(models.Product).filter(models.Product.is_deleted == False)  # noqa: E712
    if category:
        query = query.filter(models.Product.category == category)
    if q:
        query = query.filter(models.Product.name.ilike(f"%{q}%"))

    visible_ids = get_visible_product_ids(x_allowed_projects)
    if visible_ids is not None:
        query = query.filter(models.Product.id.in_(visible_ids))

    return _attach_sub_item_counts(db, query.order_by(models.Product.created_at.desc()).all())


@router.get("/bulk-import-template")
def download_product_bulk_import_template():
    """A blank spreadsheet with the right headers (plus one example row)
    for bulk-importing products — matches what /bulk-import expects."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(["Name", "Category", "Price", "Description"])
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    ws.append(["Dell Vostro 15", "Computer", 950.00, "15-inch business laptop, Intel i5, 8GB RAM, 512GB SSD"])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 55

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="product-bulk-import-template.xlsx"'},
    )


@router.post("/bulk-import")
def bulk_import_products(
    excel_file: UploadFile = File(...),
    catalogue_zip: UploadFile | None = File(None),
    db: Session = Depends(get_db),
):
    """
    Creates many products at once from a spreadsheet, optionally matched
    up with catalogue files from a zip.

    Excel file: first row is headers (case-insensitive) — needs a "Name"
    column; "Category", "Price", and "Description" are optional. Item
    Number/SKU is always auto-generated from the category, same as
    creating one product at a time.

    Catalogue archive (optional, .zip or .rar): files named "1.ext",
    "2.ext", "3.ext"…
    (any extension, subfolders ignored) — matched by POSITION to the
    spreadsheet's data rows: file "1" attaches to the first product row,
    "2" to the second, and so on. A blank/skipped row still "uses up"
    its number, so numbering always matches the row's position in the
    sheet, not the count of products actually created.
    """
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(excel_file.file.read()), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="couldn't read that file — make sure it's a valid .xlsx")

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="that spreadsheet is empty")

    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    def col_index(*names):
        for name in names:
            if name in header:
                return header.index(name)
        return None

    name_idx = col_index("name")
    category_idx = col_index("category")
    price_idx = col_index("price", "unit_price", "unit price")
    description_idx = col_index("description")

    if name_idx is None:
        raise HTTPException(status_code=400, detail='the spreadsheet needs a "Name" column')

    def cell(row, idx):
        if idx is None or idx >= len(row):
            return None
        value = row[idx]
        if value is None or (isinstance(value, str) and not value.strip()):
            return None
        return value

    # Pull every "N.ext" file out of the archive, keyed by its position number.
    files_by_position = {}
    if catalogue_zip is not None:
        entries = _extract_archive_entries(catalogue_zip.filename, catalogue_zip.file.read())
        for basename, data in entries:
            stem = basename.rsplit(".", 1)[0] if "." in basename else basename
            if stem.isdigit():
                files_by_position[int(stem)] = (basename, data)

    created = []
    for position, row in enumerate(rows[1:], start=1):
        name = cell(row, name_idx)
        if not name:
            continue  # blank row — skipped, but still "uses up" this position number
        name = str(name).strip()

        category = cell(row, category_idx)
        category = str(category).strip() if category else None

        price = None
        raw_price = cell(row, price_idx)
        if raw_price is not None:
            try:
                price = Decimal(str(raw_price))
            except InvalidOperation:
                price = None

        description = cell(row, description_idx)
        description = str(description).strip() if description else None

        sku = _generate_sku(db, category)
        product = models.Product(sku=sku, name=name, category=category, unit_price=price, description=description)
        db.add(product)
        db.commit()
        db.refresh(product)
        index_product(product)

        entry = {"sku": sku, "name": name, "row": position, "has_file": False}

        if position in files_by_position:
            filename, file_bytes = files_by_position.pop(position)
            object_key = f"{sku}/{filename}"
            upload_file(io.BytesIO(file_bytes), object_key, content_type="application/octet-stream")
            product.image_object_key = object_key
            db.commit()
            entry["has_file"] = True

        created.append(entry)

    warnings = []
    if files_by_position:
        leftover = ", ".join(str(k) for k in sorted(files_by_position.keys()))
        warnings.append(f"{len(files_by_position)} catalogue file(s) didn't match any spreadsheet row (numbers: {leftover})")

    return {"created": created, "warnings": warnings}


@router.get("/trash", response_model=list[schemas.ProductOut])
def list_trashed_products(
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    """The recycle bin — soft-deleted products, same project-visibility
    rule as the main list."""
    query = db.query(models.Product).filter(models.Product.is_deleted == True)  # noqa: E712

    visible_ids = get_visible_product_ids(x_allowed_projects)
    if visible_ids is not None:
        query = query.filter(models.Product.id.in_(visible_ids))

    return _attach_sub_item_counts(db, query.order_by(models.Product.updated_at.desc()).all())


def _check_product_visible(product, x_allowed_projects: str | None):
    visible_ids = get_visible_product_ids(x_allowed_projects)
    if visible_ids is not None and str(product.id) not in visible_ids:
        raise HTTPException(status_code=403, detail="you don't have access to this product")


@router.patch("/{product_id}/trash", response_model=schemas.ProductOut)
def trash_product(
    product_id: uuid.UUID,
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    """Soft delete — hides it from the catalogue and main list, but keeps
    it recoverable via the recycle bin. For a permanent purge, use
    DELETE /{product_id} instead (only called from within the recycle bin)."""
    product = db.query(models.Product).filter_by(id=product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="product not found")
    _check_product_visible(product, x_allowed_projects)
    product.is_deleted = True
    db.commit()
    db.refresh(product)
    return _attach_sub_item_count(db, product)


@router.patch("/{product_id}/restore", response_model=schemas.ProductOut)
def restore_product(
    product_id: uuid.UUID,
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    product = db.query(models.Product).filter_by(id=product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="product not found")
    _check_product_visible(product, x_allowed_projects)
    product.is_deleted = False
    db.commit()
    db.refresh(product)
    return _attach_sub_item_count(db, product)


@router.get("/{product_id}", response_model=schemas.ProductOut)
def get_product(
    product_id: uuid.UUID,
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    product = db.query(models.Product).filter_by(id=product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="product not found")

    visible_ids = get_visible_product_ids(x_allowed_projects)
    if visible_ids is not None and str(product.id) not in visible_ids:
        raise HTTPException(status_code=403, detail="you don't have access to this product")

    return _attach_sub_item_count(db, product)


@router.patch("/{product_id}", response_model=schemas.ProductOut)
def update_product(
    product_id: uuid.UUID,
    payload: schemas.ProductUpdate,
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    product = db.query(models.Product).filter_by(id=product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="product not found")
    _check_product_visible(product, x_allowed_projects)

    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(product, field, value)

    db.commit()
    db.refresh(product)
    index_product(product)
    return _attach_sub_item_count(db, product)


@router.delete("/{product_id}", status_code=204)
def delete_product(
    product_id: uuid.UUID,
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    product = db.query(models.Product).filter_by(id=product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="product not found")
    _check_product_visible(product, x_allowed_projects)
    db.delete(product)
    db.commit()
    remove_product_from_index(product_id)


@router.post("/{product_id}/file")
def upload_product_file(
    product_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    product = db.query(models.Product).filter_by(id=product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="product not found")
    _check_product_visible(product, x_allowed_projects)

    object_key = f"{product.sku}/{file.filename}"
    upload_file(file.file, object_key, content_type=file.content_type or "application/octet-stream")

    product.image_object_key = object_key
    db.commit()
    return {"image_object_key": object_key}


@router.get("/{product_id}/file-url")
def get_product_file_url(
    product_id: uuid.UUID,
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    product = db.query(models.Product).filter_by(id=product_id).first()
    if not product or not product.image_object_key:
        raise HTTPException(status_code=404, detail="no file attached to this product")
    _check_product_visible(product, x_allowed_projects)
    return {"url": get_presigned_url(product.image_object_key)}


@router.get("/{product_id}/file-content")
def get_product_file_content(
    product_id: uuid.UUID,
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    """
    Called by document-service (internally, no X-Allowed-Projects header —
    treated as unrestricted, same as every other internal service-to-
    service call here) to pull a product's raw catalogue file bytes when
    building a per-document catalogue zip. Also reachable directly through
    nginx's /api/products routing, which DOES forward the header, so it's
    still project-gated for that path. Not meant for direct browser use:
    no presigned convenience, no download filename handling — the
    browser-facing equivalent is /file-url.
    """
    product = db.query(models.Product).filter_by(id=product_id).first()
    if not product or not product.image_object_key:
        raise HTTPException(status_code=404, detail="no file attached to this product")
    _check_product_visible(product, x_allowed_projects)
    try:
        file_bytes = download_file_bytes(product.image_object_key)
    except Exception:
        raise HTTPException(status_code=404, detail="the file for this product is missing from storage")
    return Response(content=file_bytes, media_type="application/octet-stream")


# ---------- sub-items (components a "bundle" product is made of) ----------

def _renumber_sub_items(db: Session, parent_product_id: uuid.UUID):
    """Keeps sequence numbers contiguous (1, 2, 3…) after a removal —
    otherwise a deleted #2 would leave a gap and #3 would stay "3" forever,
    which would be a confusing mismatch with the zip's filenames."""
    remaining = (
        db.query(models.ProductSubItem)
        .filter_by(parent_product_id=parent_product_id)
        .order_by(models.ProductSubItem.sequence_number)
        .all()
    )
    for i, item in enumerate(remaining, start=1):
        item.sequence_number = i
    db.commit()


@router.get("/{product_id}/sub-items", response_model=list[schemas.SubItemOut])
def list_sub_items(product_id: uuid.UUID, db: Session = Depends(get_db)):
    rows = (
        db.query(models.ProductSubItem)
        .filter_by(parent_product_id=product_id)
        .order_by(models.ProductSubItem.sequence_number)
        .all()
    )
    result = []
    for row in rows:
        sub_product = db.query(models.Product).filter_by(id=row.sub_product_id).first()
        if not sub_product:
            continue  # the referenced product was hard-deleted since — skip it rather than error
        result.append(schemas.SubItemOut(
            id=row.id,
            sequence_number=row.sequence_number,
            product_id=sub_product.id,
            sku=sub_product.sku,
            name=sub_product.name,
            description=sub_product.description,
            unit_price=sub_product.unit_price,
            currency=sub_product.currency,
            has_file=bool(sub_product.image_object_key),
        ))
    return result


@router.post("/{product_id}/sub-items", response_model=schemas.SubItemOut, status_code=201)
def add_sub_item(product_id: uuid.UUID, payload: schemas.SubItemAdd, db: Session = Depends(get_db)):
    parent = db.query(models.Product).filter_by(id=product_id).first()
    if not parent:
        raise HTTPException(status_code=404, detail="product not found")

    if payload.product_id == product_id:
        raise HTTPException(status_code=400, detail="a product can't be its own sub-item")

    sub_product = db.query(models.Product).filter_by(id=payload.product_id).first()
    if not sub_product:
        raise HTTPException(status_code=404, detail="that product doesn't exist")

    next_seq = (
        db.query(func.max(models.ProductSubItem.sequence_number))
        .filter_by(parent_product_id=product_id)
        .scalar() or 0
    ) + 1

    sub_item = models.ProductSubItem(
        parent_product_id=product_id,
        sub_product_id=payload.product_id,
        sequence_number=next_seq,
    )
    db.add(sub_item)
    db.commit()
    db.refresh(sub_item)

    return schemas.SubItemOut(
        id=sub_item.id,
        sequence_number=sub_item.sequence_number,
        product_id=sub_product.id,
        sku=sub_product.sku,
        name=sub_product.name,
        description=sub_product.description,
        unit_price=sub_product.unit_price,
        currency=sub_product.currency,
        has_file=bool(sub_product.image_object_key),
    )


@router.delete("/{product_id}/sub-items/{sub_item_id}", status_code=204)
def remove_sub_item(product_id: uuid.UUID, sub_item_id: uuid.UUID, db: Session = Depends(get_db)):
    sub_item = (
        db.query(models.ProductSubItem)
        .filter_by(id=sub_item_id, parent_product_id=product_id)
        .first()
    )
    if not sub_item:
        raise HTTPException(status_code=404, detail="sub-item not found")
    db.delete(sub_item)
    db.commit()
    _renumber_sub_items(db, product_id)


@router.get("/{product_id}/download-bundle")
def download_product_bundle(
    product_id: uuid.UUID,
    db: Session = Depends(get_db),
    x_allowed_projects: str | None = Header(default=None, alias="X-Allowed-Projects"),
):
    """
    For a product with sub-items: fetches every sub-item's catalogue file,
    renames each to its sequence number (1, 2, 3… keeping the original
    extension), and streams back a single zip. For a product with no
    sub-items, the frontend should use /file-url instead — this endpoint
    only makes sense once there's something to bundle.
    """
    product = db.query(models.Product).filter_by(id=product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="product not found")
    _check_product_visible(product, x_allowed_projects)

    sub_items = (
        db.query(models.ProductSubItem)
        .filter_by(parent_product_id=product_id)
        .order_by(models.ProductSubItem.sequence_number)
        .all()
    )
    if not sub_items:
        raise HTTPException(status_code=404, detail="this product has no sub-items to bundle")

    zip_buffer = io.BytesIO()
    added_any = False
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for item in sub_items:
            sub_product = db.query(models.Product).filter_by(id=item.sub_product_id).first()
            if not sub_product or not sub_product.image_object_key:
                continue  # nothing to include for this one — skip rather than fail the whole zip
            ext = ""
            if "." in sub_product.image_object_key.rsplit("/", 1)[-1]:
                ext = "." + sub_product.image_object_key.rsplit(".", 1)[-1]
            try:
                file_bytes = download_file_bytes(sub_product.image_object_key)
            except Exception:
                continue  # object missing from storage — skip rather than fail the whole zip
            zf.writestr(f"{item.sequence_number}{ext}", file_bytes)
            added_any = True

    if not added_any:
        raise HTTPException(status_code=404, detail="none of this product's sub-items have a catalogue file")

    zip_buffer.seek(0)
    filename = f"{product.sku}-catalogue.zip"
    return Response(
        content=zip_buffer.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
