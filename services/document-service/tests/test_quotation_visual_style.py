"""
Regression tests for the quotation's visual styling: heading size/colour,
page margins and vertical spacing, logo size, and the Customer Service
footer text. Companion to test_quotation_pdf_layout.py (pagination) — same
approach: call generate_quotation_pdf directly with lightweight fake
objects, no Postgres/Docker needed.

Unlike the pagination tests (which only read text out of the finished PDF
with pypdf), these need the *rendered CSS values* — font size, colour,
box size, vertical position — which a PDF text extract can't give back
reliably (fonts are subset-embedded, colours live in graphics-state
operators). So `layout` below captures the exact HTML that
generate_quotation_pdf hands to WeasyPrint, re-renders it with WeasyPrint's
own layout engine, and walks the resulting box tree: the same numbers the
PDF is drawn from, with no new dependency.

All sizes WeasyPrint reports are CSS px (96 per inch): 1pt = 4/3 px.
"""
import io
from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from PIL import Image
from weasyprint import HTML as RealHTML

from app.core import export_pdf
from app.core.export_pdf import generate_quotation_pdf
from app.core.export_quotation import generate_quotation_xlsx

PT = 96 / 72  # px per pt

COMPANY = {
    "name": "Acme Supplier Co., Ltd.",
    "position": None,
    "address": "No. 1, Sample Road, Yangon",
    "contact_no": "+95 1 234 5678",
    "support_email": "sales@acme.example",
}

CUSTOMER = {
    "id": "11111111-1111-1111-1111-111111111111",
    "name": "Beta Customer Co., Ltd.",
    "billing_address": {"line1": "No. 99, Customer Street", "city": "Mandalay"},
}

DIRECTOR = {
    "name": "Zaw Zaw",
    "address": "N/A",
    "contact_no": "N/A",
    "email": "director@acme.example",
}


def _png(size=(200, 200), color="navy"):
    """A 200x200 PNG — the size the app standardizes every uploaded logo/seal to."""
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def _document():
    return SimpleNamespace(
        issue_date=date(2026, 1, 1),
        currency="USD",
        doc_number="QT-0001",
        terms_and_conditions=None,
        subtotal=None,
        tax_total=None,
        total=None,
    )


def _line_item(i):
    return SimpleNamespace(
        description=f"Item {i} with a reasonably long name for the row",
        quantity=1,
        unit_price=100,
        unit="Nos",
        remark="",
        product_id=None,
        tax_rate=None,
    )


def _walk(box):
    yield box
    for child in getattr(box, "children", None) or []:
        yield from _walk(child)


def _texts(boxes, needle):
    """Text boxes whose text contains `needle`."""
    return [b for b in boxes if getattr(b, "text", None) and needle in b.text]


def _images(boxes):
    return [b for b in boxes if type(b).__name__ in ("InlineReplacedBox", "BlockReplacedBox")]


@pytest.fixture
def layout(monkeypatch):
    """
    layout(n_items=3, **generate_quotation_pdf kwargs) -> SimpleNamespace(
        pdf=<BytesIO>, html=<str>, pages=[[boxes of page 1], [boxes of page 2], ...]).
    """

    def build(n_items=3, company=COMPANY, customer=CUSTOMER, director=DIRECTOR, **kwargs):
        captured = {}

        class _Spy:
            def __init__(self, string=None, **kw):
                captured["html"] = string
                self._real = RealHTML(string=string, **kw)

            def write_pdf(self, *args, **kw):
                return self._real.write_pdf(*args, **kw)

        monkeypatch.setattr(export_pdf, "HTML", _Spy)
        items = [(_line_item(i), None, []) for i in range(1, n_items + 1)]
        pdf = generate_quotation_pdf(
            _document(), customer, items, company=company, director=director, **kwargs
        )
        assert pdf.getvalue()[:4] == b"%PDF"
        rendered = RealHTML(string=captured["html"]).render()
        return SimpleNamespace(
            pdf=pdf,
            html=captured["html"],
            pages=[list(_walk(p._page_box)) for p in rendered.pages],
        )

    return build


# ---------------------------------------------------------------------------
# Task 1 — QUOTATION heading: smaller and red
# ---------------------------------------------------------------------------

OLD_HEADING_PX = 22 * PT  # what the heading used to be (22pt)
BODY_PX = 10.5 * PT


def _heading(layout_result):
    boxes = _texts(layout_result.pages[0], "QUOTATION")
    assert len(boxes) == 1, "the QUOTATION heading must render exactly once on page 1"
    return boxes[0]


def test_heading_is_red(layout):
    red, green, blue, alpha = _heading(layout()).style["color"]
    assert red > 0.6 and green < 0.1 and blue < 0.1 and alpha == 1, (
        f"QUOTATION heading must be red, got rgba({red:.2f}, {green:.2f}, {blue:.2f}, {alpha})"
    )


def test_heading_is_same_size_as_the_date_line(layout):
    result = layout()
    size = _heading(result).style["font_size"]
    date_size = _texts(result.pages[0], "Date:")[0].style["font_size"]
    assert size == pytest.approx(date_size, abs=0.01), f"heading {size / PT:.1f}pt != date {date_size / PT:.1f}pt"
    assert size < OLD_HEADING_PX


def test_heading_stays_bold_and_left_aligned_above_the_date(layout):
    result = layout()
    heading = _heading(result)
    assert heading.style["font_weight"] >= 700
    date_line = _texts(result.pages[0], "Date:")[0]
    supplier = _texts(result.pages[0], "SUPPLIER")[0]
    assert heading.position_y < date_line.position_y < supplier.position_y
    assert heading.position_x == pytest.approx(supplier.position_x, abs=1), (
        "heading must stay on the same left edge as the SUPPLIER block"
    )


def test_other_headings_and_body_text_are_not_affected_by_the_heading_change(layout):
    result = layout()
    for needle in ("SUPPLIER", "END USER", "Terms and Conditions"):
        box = _texts(result.pages[0] + (result.pages[1] if len(result.pages) > 1 else []), needle)[0]
        assert box.style["font_size"] == pytest.approx(11 * PT, abs=0.01), f"{needle} heading size changed"
        assert box.style["color"][:3] == pytest.approx((0x1A / 255,) * 3, abs=0.01), f"{needle} heading colour changed"
    body = _texts(result.pages[0], "Sample Road")[0]
    assert body.style["font_size"] == pytest.approx(10 * PT, abs=0.01)
    assert body.style["color"][:3] == pytest.approx((0x1A / 255,) * 3, abs=0.01)


def test_xlsx_heading_is_red_and_smaller_too():
    wb = load_workbook(
        generate_quotation_xlsx(_document(), CUSTOMER, [(_line_item(1), None, [])], company=COMPANY, director=DIRECTOR)
    )
    cells = [c for row in wb.active.iter_rows() for c in row if c.value == "QUOTATION"]
    assert len(cells) == 1
    font = cells[0].font
    assert font.color is not None and font.color.rgb.endswith("C00000"), "XLSX heading must be red"
    date_font = next(c for row in wb.active.iter_rows() for c in row if c.value == "Date:").font
    assert font.sz == (date_font.sz or 11), "XLSX heading must match the Date line size"
    assert font.b


# ---------------------------------------------------------------------------
# Task 2 — page margins and vertical whitespace
# ---------------------------------------------------------------------------

# Measured on the original template (WeasyPrint box tree, CSS px):
#   first content 64.7 (page margin 56.7 + browser-default 8px body margin),
#   logo-row bottom -> heading 29, parties -> table 18, 100px min-height
#   header row even with no logo.
OLD_TOP_CONTENT_PX = 64.7
MAX_TOP_CONTENT_PX = 50  # 1.2cm page margin (45.4px) + a little slack


def _page_box(layout_result):
    return layout_result.pages[0][0]


def _first_content_y(layout_result):
    boxes = [b for b in layout_result.pages[0] if getattr(b, "text", None) or _images([b])]
    return min(b.position_y for b in boxes)


def test_top_of_first_page_has_no_stacked_margins(layout):
    result = layout(logo_bytes=_png(), seal_bytes=_png())
    top = _first_content_y(result)
    assert top < OLD_TOP_CONTENT_PX - 8, f"first content starts {top:.1f}px from the page edge, expected a smaller top margin"
    assert top <= MAX_TOP_CONTENT_PX
    assert top >= 0.8 * 96 / 2.54, "top margin must not become so small that printers clip the logo"


def test_gap_between_logo_row_and_heading_is_tight(layout):
    result = layout(logo_bytes=_png(), seal_bytes=_png())
    heading = _texts(result.pages[0], "QUOTATION")[0]
    imgs = _images(result.pages[0])
    assert imgs
    logo_bottom = max(i.position_y + i.height for i in imgs[:2])
    gap = heading.position_y - logo_bottom
    assert 0 <= gap <= 16, f"gap between logo row and heading is {gap:.1f}px (was 29)"


def test_no_blank_header_row_when_company_has_no_logo(layout):
    result = layout()
    heading = _texts(result.pages[0], "QUOTATION")[0]
    assert heading.position_y < 80, f"heading sits {heading.position_y:.1f}px down with no logo (a blank 100px row was reserved before)"


def test_heading_date_and_parties_spacing_is_compact_but_separated(layout):
    result = layout(logo_bytes=_png(), seal_bytes=_png())
    p = result.pages[0]
    heading, date_line = _texts(p, "QUOTATION")[0], _texts(p, "Date:")[0]
    supplier, end_user = _texts(p, "SUPPLIER")[0], _texts(p, "END USER")[0]
    assert date_line.position_y - (heading.position_y + heading.height) <= 10
    assert 4 <= supplier.position_y - (date_line.position_y + date_line.height) <= 20
    assert supplier.position_y == pytest.approx(end_user.position_y, abs=0.5), "SUPPLIER / END USER headings must stay aligned"


def test_items_table_starts_close_below_the_parties_block(layout):
    result = layout(n_items=3, logo_bytes=_png(), seal_bytes=_png())
    p = result.pages[0]
    last_party_row = max(b.position_y + b.height for b in _texts(p, "Address") + _texts(p, "Customer Street"))
    header = _texts(p, "Price (USD)")[0]
    gap = header.position_y - last_party_row
    assert 0 < gap <= 30, f"{gap:.1f}px between the last SUPPLIER/END USER row and the table header (was ~40+)"


def test_long_wrapped_address_still_does_not_overlap_the_table(layout):
    company = dict(COMPANY, address="Building 12, Floor 7, Room 704, Very Long Road Name, Kamayut Township, Yangon Region, Myanmar " * 2)
    result = layout(company=company, logo_bytes=_png(), seal_bytes=_png())
    p = result.pages[0]
    addr = _texts(p, "Very Long Road")
    bottom = max(b.position_y + b.height for b in addr)
    assert bottom < _texts(p, "Price (USD)")[0].position_y


def test_multi_page_quotation_keeps_margins_and_content_inside_the_page(layout):
    result = layout(n_items=40, logo_bytes=_png(), seal_bytes=_png())
    assert len(result.pages) >= 2
    page_w, page_h = 793.7, 1122.5  # A4 in CSS px
    for n, boxes in enumerate(result.pages, 1):
        for b in boxes:
            if getattr(b, "text", None):
                assert b.position_x >= 0.5 * 96 / 2.54 and b.position_x + b.width <= page_w - 0.5 * 96 / 2.54, f"page {n}: '{b.text[:20]}' outside side margins"
                assert b.position_y + b.height <= page_h, f"page {n}: '{b.text[:20]}' runs off the page"
    # later pages start at the same compact top margin (no header repeated there)
    assert _first_content_y(SimpleNamespace(pages=[result.pages[1]])) <= MAX_TOP_CONTENT_PX + 25


# ---------------------------------------------------------------------------
# Task 4 — header logo (and company seal) about 10px smaller, ratio preserved
# ---------------------------------------------------------------------------

def _header_images(layout_result):
    imgs = sorted(_images(layout_result.pages[0]), key=lambda i: i.position_x)
    return imgs


def test_header_logo_and_seal_are_120px_down_from_130(layout):
    imgs = _header_images(layout(logo_bytes=_png(), seal_bytes=_png()))
    assert len(imgs) == 2
    for img in imgs:
        assert img.width == pytest.approx(120, abs=0.5) and img.height == pytest.approx(120, abs=0.5), (
            f"header image is {img.width:.0f}x{img.height:.0f}px, expected 120x120 (was 130x130)")


def test_logo_aspect_ratio_is_preserved_for_non_square_images(layout):
    wide = _header_images(layout(logo_bytes=_png((400, 200))))[0]
    tall = _header_images(layout(logo_bytes=_png((200, 400))))[0]
    assert wide.width / wide.height == pytest.approx(2.0, rel=0.02) and wide.width <= 120.5
    assert tall.width / tall.height == pytest.approx(0.5, rel=0.02) and tall.height <= 120.5


def test_small_logo_is_not_stretched(layout):
    img = _header_images(layout(logo_bytes=_png((60, 60))))[0]
    assert img.width == pytest.approx(60, abs=0.5) and img.height == pytest.approx(60, abs=0.5)


def test_logo_stays_left_and_seal_right_inside_margins_and_clear_of_heading(layout):
    result = layout(logo_bytes=_png(), seal_bytes=_png())
    logo, seal = _header_images(result)
    heading = _texts(result.pages[0], "QUOTATION")[0]
    supplier = _texts(result.pages[0], "SUPPLIER")[0]
    assert logo.position_x == pytest.approx(supplier.position_x, abs=1), "logo must stay on the left content edge"
    assert seal.position_x + seal.width <= 793.7 - 2 * 96 / 2.54 + 1, "seal must stay inside the right margin"
    assert logo.position_y + logo.height <= heading.position_y, "logo overlaps the heading"
    assert logo.position_y == pytest.approx(seal.position_y, abs=0.5), "logo and seal must share a top edge"


def test_logo_only_on_first_page_and_md_seal_size_unchanged(layout):
    result = layout(n_items=40, logo_bytes=_png(), seal_bytes=_png(), director_seal_bytes=_png())
    assert len(result.pages) >= 2
    assert len(_images(result.pages[0])) == 2
    assert not any(_images(p) for p in result.pages[1:-1])  # no stray header images on middle pages
    md = [i for p in result.pages for i in _images(p) if i.width == pytest.approx(130, abs=0.5)]
    assert md, "the MD signature seal keeps its 130px size (not a logo)"


def test_xlsx_logo_and_seal_are_10px_smaller():
    wb = load_workbook(generate_quotation_xlsx(
        _document(), CUSTOMER, [(_line_item(1), None, [])], company=COMPANY,
        logo_bytes=_png(), seal_bytes=_png(), director=None))
    imgs = wb.active._images
    assert len(imgs) == 2
    # A reloaded image reports its pixel size; the displayed size is the anchor extent (EMU, 9525 per px).
    assert all(i.anchor.ext.width == 90 * 9525 and i.anchor.ext.height == 90 * 9525 for i in imgs)


# ---------------------------------------------------------------------------
# Task 5 — Customer Service footer: smaller Email / Phone, same heading
# ---------------------------------------------------------------------------

def _footer_boxes(result, needle):
    out = []
    for page in result.pages:
        out += _texts(page, needle)
    return out


def test_footer_email_and_phone_are_smaller_than_body_text(layout):
    result = layout()
    for needle in ("Email: sales@acme.example", "Phone: +95 1 234 5678"):
        boxes = _footer_boxes(result, needle)
        assert boxes, f"{needle!r} not rendered in the footer"
        for b in boxes:
            assert b.style["font_size"] < 10.5 * PT, f"{needle} is {b.style['font_size'] / PT:.1f}pt, not smaller than the 10.5pt it was"
            assert b.style["font_size"] >= 7 * PT, "footer text must stay legible"


def test_footer_heading_style_is_unchanged(layout):
    for b in _footer_boxes(layout(), "Customer Service"):
        assert b.style["font_size"] == pytest.approx(11 * PT, abs=0.01)
        assert b.style["font_weight"] >= 700
        assert b.style["color"][:3] == pytest.approx((0x1A / 255,) * 3, abs=0.01)


def test_footer_repeats_on_every_page_and_is_inside_the_page(layout):
    result = layout(n_items=40)
    assert len(result.pages) >= 2
    for n, page in enumerate(result.pages, 1):
        email = _texts(page, "Email: sales@acme.example")
        assert email, f"page {n} lost its footer"
        assert email[0].position_y + email[0].height <= 1122.5 - 0.5 * 96 / 2.54, f"page {n} footer is clipped by the page bottom"


def test_long_footer_values_wrap_instead_of_being_clipped(layout):
    company = dict(COMPANY, support_email="very.long.customer.service.mailbox.name.for.testing@a-really-long-company-domain-name.example.com",
                   contact_no="+95 1 234 5678, +95 9 111 222 333, +95 9 444 555 666, +95 9 777 888 999")
    result = layout(company=company, n_items=40)
    page_right = 793.7 - 2 * 96 / 2.54
    for n, page in enumerate(result.pages, 1):
        lines = [b for b in page if getattr(b, "text", None) and b.position_y > 900]
        assert lines
        for b in lines:
            assert b.position_x + b.width <= page_right + 1, f"page {n}: footer text '{b.text[:25]}' runs past the right margin"
            assert b.position_y + b.height <= 1122.5, f"page {n}: footer text clipped at the bottom"
    text = " ".join(b.text for b in result.pages[0] if getattr(b, "text", None))
    assert "very.long.customer" in text and "+95 9 777" in text or "777" in text


def test_long_unbreakable_supplier_values_wrap_inside_their_column(layout):
    company = dict(COMPANY, support_email="x" * 70 + "@" + "y" * 40 + ".example.com")
    director = dict(DIRECTOR, position="Chief-Executive-and-Managing-Director-of-Group-Operations-Division")
    result = layout(company=company, director=director)
    p = result.pages[0]
    end_user_x = _texts(p, "END USER")[0].position_x
    for needle in ("xxxxx", "Chief-Executive", "Acme"):
        for b in [b for b in _texts(p, needle) if b.position_y < 600]:  # parties block, not the footer
            assert b.position_x + b.width <= end_user_x, f"'{b.text[:20]}' spills into the END USER column"
