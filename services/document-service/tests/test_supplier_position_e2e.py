"""
End-to-end check of the Supplier "Position" row through the REAL HTTP API
and ORM (SQLite stands in for Postgres; JSONB is mapped to JSON there):

  PATCH/POST /companies/{id}/directors  ->  POST /documents (director_id)
  ->  GET /documents/{id}/export/quotation-pdf  and  .../quotation

Only the outside world is stubbed (search indexing, audit log).
"""
import io
import uuid
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.db import Base, get_db
from app.main import app
from app.routers import documents as documents_router


@compiles(JSONB, "sqlite")
def _jsonb_as_json(type_, compiler, **kw):
    return "JSON"


@pytest.fixture
def client(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False)

    def _get_db():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = _get_db
    monkeypatch.setattr(documents_router, "index_document", lambda doc: None)
    monkeypatch.setattr(documents_router, "log_action", lambda *a, **k: None)
    # no `with`: skips startup (migrations/Postgres). The gateway secret is the real one nginx would add.
    yield TestClient(app, headers={"X-Internal-Secret": settings.internal_shared_secret})
    app.dependency_overrides.clear()


def _setup(client, company_position=None):
    company = client.post("/companies", json={"name": "Aungaung", "short_name": "AA", "position": company_position,
                                              "address": "Mayangon", "contact_no": "+9599", "support_email": "a@b.example"}).json()
    return company["id"]


def _director(client, company_id, name, position):
    body = {"name": name}
    if position is not None:
        body["position"] = position
    r = client.post(f"/companies/{company_id}/directors", json=body)
    assert r.status_code == 201, r.text
    # iteration 12: an MD's position is required; a missing one defaults to "Managing Director"
    assert r.json()["position"] == (position if position is not None else "Managing Director")
    return r.json()["id"]


def _quotation(client, company_id, director_id):
    r = client.post("/documents", json={"doc_type": "quotation", "company_id": company_id, "director_id": director_id,
                                        "issue_date": str(date(2026, 10, 9)), "items": []})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _supplier_block(client, doc_id):
    r = client.get(f"/documents/{doc_id}/export/quotation-pdf")
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
    return text.split("SUPPLIER", 1)[1].split("END USER", 1)[0]


def _xlsx_rows(client, doc_id):
    r = client.get(f"/documents/{doc_id}/export/quotation")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    return [[c.value for c in row] for row in ws.iter_rows() if any(c.value is not None for c in row)]


def test_selected_directors_position_reaches_pdf_and_excel_in_the_right_place(client):
    cid = _setup(client)
    md = _director(client, cid, "Aung Aung", "MD")
    ceo = _director(client, cid, "Zaw Zaw", "Chief Executive Officer")
    plain = _director(client, cid, "No Title", None)

    pdf = _supplier_block(client, _quotation(client, cid, md))
    assert pdf.index("Name") < pdf.index("Position") < pdf.index("Address")
    assert "MD" in pdf.split("Position", 1)[1].split("Address", 1)[0]

    other = _supplier_block(client, _quotation(client, cid, ceo))
    assert "Chief Executive Officer" in other and "MD" not in other.split("Position", 1)[1].split("Address", 1)[0]

    # iteration 12: an MD created without a position gets "Managing Director"
    assert "Managing Director" in _supplier_block(client, _quotation(client, cid, plain))
    assert "Position" not in _supplier_block(client, _quotation(client, cid, None))   # legacy: no MD at all

    rows = _xlsx_rows(client, _quotation(client, cid, md))
    pos = next(r for r in rows if r[0] == "Position")
    assert pos[1] == "MD"
    labels = [r[0] for r in rows]
    assert labels.index("SUPPLIER: Aungaung") < labels.index("Position") < labels.index("Address")  # company name now on the heading
    plain_rows = _xlsx_rows(client, _quotation(client, cid, plain))
    assert next(r for r in plain_rows if r[0] == "Position")[1] == "Managing Director"


def test_editing_a_directors_position_changes_the_next_export(client):
    cid = _setup(client)
    did = _director(client, cid, "Aung Aung", "Director")
    doc = _quotation(client, cid, did)
    assert "Director" in _supplier_block(client, doc)
    r = client.patch(f"/companies/{cid}/directors/{did}", json={"position": "MD"})
    assert r.status_code == 200 and r.json()["position"] == "MD"
    block = _supplier_block(client, doc)  # same document, position is looked up live
    assert "MD" in block and "Director" not in block
    # iteration 12: the position can't be cleared (422) and the export is unchanged
    assert client.patch(f"/companies/{cid}/directors/{did}", json={"position": None}).status_code == 422
    assert client.patch(f"/companies/{cid}/directors/{did}", json={"position": "  "}).status_code == 422
    assert "MD" in _supplier_block(client, doc)


def test_company_signer_position_is_no_longer_used(client):
    """Inverted in iteration 12 (was ..._is_the_fallback_when_the_md_has_none):
    a company `position` sent by an old client is ignored and never printed."""
    cid = _setup(client, company_position="Director")
    plain = _director(client, cid, "No Title", None)
    md = _director(client, cid, "Aung Aung", "MD")
    assert "Managing Director" in _supplier_block(client, _quotation(client, cid, plain))
    assert "Position" not in _supplier_block(client, _quotation(client, cid, None))  # no MD -> no Position row
    own = _supplier_block(client, _quotation(client, cid, md))
    assert "MD" in own and "Director" not in own
