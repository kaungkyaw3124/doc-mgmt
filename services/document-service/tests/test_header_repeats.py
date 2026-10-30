"""
Loop iteration 2 — the quotation header (logo + seal, QUOTATION, Date,
SUPPLIER / END USER blocks) repeats at the top of EVERY PDF page, followed
by the items table's column titles where the table continues. This
reverses the earlier "header on page 1 only" rule; see the comments in
quotation.html. It is done with the <thead> of an outer layout table (so
its height is the header's real height), never a fixed-height `running()`
margin box, so the checks below include a long, wrapping address.

The Excel export can only repeat rows when printed (print titles), so its
check is that ws.print_title_rows covers row 1 through the column titles.

Same approach as test_quotation_visual_style.py: generate_quotation_pdf is
called directly with lightweight fakes; text is read back with pypdf, and
geometry from WeasyPrint's own box tree (via that module's `layout`
fixture, imported here).
"""
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.export_quotation import generate_quotation_xlsx
from tests.test_quotation_visual_style import (  # noqa: F401  (layout is a fixture)
    COMPANY,
    CUSTOMER,
    DIRECTOR,
    _document,
    _images,
    _line_item,
    _png,
    _texts,
    layout,
)

HEADER_FIELDS = ("QUOTATION", "Date:", f"SUPPLIER: {COMPANY['name']}", "END USER",
                 f"Name {DIRECTOR['name']}", "QT-0001")
CUSTOMER_ADDRESS_LINE = "Address No. 99, Customer Street, Mandalay"  # last header line in pypdf's text


def _pages_text(result):
    return [p.extract_text() or "" for p in PdfReader(BytesIO(result.pdf.getvalue())).pages]


def _header_bottom(page_boxes):
    headers = [b for b in page_boxes
               if getattr(b, "element", None) is not None and b.element.get("class") == "pdf-header"]
    assert len(headers) == 1, f"expected exactly one .pdf-header box on the page, got {len(headers)}"
    h = headers[0]
    return h.position_y + h.margin_height()


def test_80_items_header_on_every_page_and_column_titles_follow_it(layout):
    result = layout(n_items=80, logo_bytes=_png(), seal_bytes=_png(), director_seal_bytes=_png())
    pages = _pages_text(result)
    assert len(pages) >= 3, f"80 items should need at least 3 pages, got {len(pages)}"
    for n, text in enumerate(pages, 1):
        for field in HEADER_FIELDS:
            assert field in text, f"page {n}/{len(pages)} is missing header field {field!r}"
        if "with a reasonably long name for the row" in text:  # page has item rows
            header_end = text.index(CUSTOMER_ADDRESS_LINE) + len(CUSTOMER_ADDRESS_LINE)
            titles = text.find("No Item Qty")
            assert titles >= header_end, f"page {n}: column titles must come after the header"
            assert text.index("with a reasonably long name for the row") > titles, (
                f"page {n}: item rows must come after the column titles"
            )


def test_long_wrapping_address_header_never_overlaps_the_items(layout):
    long_address = ("Building 12, Floor 7, Room 704, Very Long Road Name, Kamayut Township, "
                    "Yangon Region, Myanmar ") * 3
    # On the director: the Supplier Address row shows the selected MD's address.
    result = layout(n_items=80, director=dict(DIRECTOR, address=long_address),
                    logo_bytes=_png(), seal_bytes=_png())
    assert len(result.pages) >= 3

    # the address really wraps to 5+ lines in the SUPPLIER column
    first = result.pages[0]
    page_mid = 793.7 / 2
    addr_lines = {round(b.position_y) for b in first
                  if getattr(b, "text", None) and b.position_x < page_mid
                  and any(w in b.text for w in ("Building", "Floor", "Road", "Township", "Region", "Myanmar"))}
    assert len(addr_lines) >= 5, f"test setup: the address only wrapped to {len(addr_lines)} line(s)"

    for n, page in enumerate(result.pages, 1):
        bottom = _header_bottom(page)
        for b in page:
            if getattr(b, "text", None) and any(w in b.text for w in ("Myanmar", "Customer Street")):
                assert b.position_y + b.height <= bottom + 0.5, f"page {n}: header text spills below the header box"
        titles = _texts(page, "Price (USD)")
        rows = [b for b in page if getattr(b, "text", None) and "with a reasonably long name" in b.text]
        if rows:
            assert titles, f"page {n}: item rows without column titles"
            assert bottom <= titles[0].position_y, f"page {n}: header ({bottom:.1f}px) overlaps the column titles"
            assert bottom < min(r.position_y for r in rows), f"page {n}: header overlaps the first item row"


@pytest.mark.parametrize("n_items", [1, 2])
def test_short_quotation_still_fits_on_one_page(layout, n_items):
    result = layout(n_items=n_items, logo_bytes=_png(), seal_bytes=_png(), director_seal_bytes=_png())
    assert len(result.pages) == 1, f"{n_items}-item quotation needs {len(result.pages)} pages, expected 1"


def test_header_without_logo_or_seal_repeats_without_gaps(layout):
    result = layout(n_items=80)
    assert len(result.pages) >= 3
    first_y = None
    for n, page in enumerate(result.pages, 1):
        assert not _images(page), f"page {n}: no images were uploaded, none should render"
        heading = _texts(page, "QUOTATION")
        assert len(heading) == 1, f"page {n}: expected the QUOTATION heading once"
        y = heading[0].position_y
        # Same bound as test_no_blank_header_row_when_company_has_no_logo
        # (page 1 only, before): a 100px+ logo row must not be reserved.
        assert y < 80, f"page {n}: heading sits {y:.1f}px down — an empty logo row is reserved"
        first_y = first_y if first_y is not None else y
        assert y == pytest.approx(first_y, abs=0.5), f"page {n}: header starts lower than on page 1"
        assert _texts(page, f"SUPPLIER: {COMPANY['name']}") and _texts(page, "END USER")


def test_xlsx_print_titles_cover_the_header_through_the_column_titles():
    items = [(_line_item(i), None, []) for i in range(1, 81)]
    ws = load_workbook(generate_quotation_xlsx(_document(), CUSTOMER, items, company=COMPANY,
                                               director=DIRECTOR, logo_bytes=_png(), seal_bytes=_png())).active
    titles_row = next(r for r in range(1, ws.max_row + 1)
                      if ws.cell(r, 1).value == "No" and ws.cell(r, 2).value == "Item")
    assert ws.print_title_rows == f"$1:${titles_row}"
    assert titles_row > 1
    header_cells = [ws.cell(r, c).value for r in range(1, titles_row) for c in (1, 5)]
    assert f"SUPPLIER: {COMPANY['name']}" in header_cells and "END USER" in header_cells
