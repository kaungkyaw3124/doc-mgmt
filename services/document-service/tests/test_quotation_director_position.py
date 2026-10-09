"""
Task 3 — the SUPPLIER block shows the position of the MD selected on the
document. The position is stored once, on CompanyDirector (reusing the
existing documents.director_id selection) — there is no documents.position
column, so it can never go out of sync with the selected MD.
"""
import ast
import pathlib
import uuid
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app import models, schemas
from app.core.export_pdf import generate_quotation_pdf
from app.core.export_quotation import generate_quotation_xlsx
from app.routers import documents as documents_router

COMPANY = {"name": "Acme Supplier Co., Ltd.", "position": "Director", "address": "No. 1 Sample Road",
           "contact_no": "+95 1 234 5678", "support_email": "sales@acme.example"}
CUSTOMER = {"id": str(uuid.uuid4()), "name": "Beta Customer", "billing_address": {"line1": "Street 9"}}


def _doc():
    return SimpleNamespace(issue_date=date(2026, 1, 1), currency="USD", doc_number="QT-1",
                           terms_and_conditions=None, subtotal=None, tax_total=None, total=None)


def _item():
    return SimpleNamespace(description="Widget", quantity=1, unit_price=10, unit="Nos", remark="",
                           product_id=None, tax_rate=None)


def _director(**kw):
    d = {"name": "Zaw Zaw", "address": "N/A", "contact_no": "N/A", "email": "d@acme.example"}
    d.update(kw)
    return d


def _pdf_text(director):
    pdf = generate_quotation_pdf(_doc(), CUSTOMER, [(_item(), None, [])], company=COMPANY, director=director)
    return "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf.getvalue())).pages)


def _supplier_block(text):
    return text.split("SUPPLIER", 1)[1].split("END USER", 1)[0]


# ---- PDF ------------------------------------------------------------------

@pytest.mark.parametrize("position", ["MD", "Managing Director", "Chief Technology Officer"])
def test_pdf_supplier_block_shows_selected_directors_position(position):
    block = _supplier_block(_pdf_text(_director(position=position)))
    assert "Position" in block and position in block


def test_pdf_position_sits_between_name_and_address():
    block = _supplier_block(_pdf_text(_director(position="MD")))
    assert block.index("Name") < block.index("Position") < block.index("Address")


def test_pdf_two_different_directors_give_two_different_positions():
    a = _supplier_block(_pdf_text(_director(name="A", position="Founder")))
    b = _supplier_block(_pdf_text(_director(name="B", position="Finance Director")))
    assert "Founder" in a and "Finance Director" not in a
    assert "Finance Director" in b and "Founder" not in b


@pytest.mark.parametrize("director", [None, _director(), _director(position=None), _director(position="")])
def test_pdf_without_a_position_omits_the_row_and_still_renders(director):
    """Legacy documents: no MD selected, or an MD created before the column
    existed (no 'position' key at all / NULL). Must render, with no
    placeholder and no hardcoded title in the Supplier block."""
    block = _supplier_block(_pdf_text(director))
    assert "Position" not in block
    assert "Managing Director" not in block and "None" not in block
    assert "Acme Supplier Co., Ltd." in block


def test_pdf_position_is_escaped():
    block = _supplier_block(_pdf_text(_director(position="<b>MD</b>")))
    assert "<b>MD</b>" in block  # rendered literally, not interpreted as markup


def test_signature_title_uses_position_and_falls_back_for_legacy_directors():
    assert "Head of Sales" in _pdf_text(_director(position="Head of Sales")).split("Date")[-1]
    assert "Managing Director" in _pdf_text(_director())  # legacy: unchanged output


# ---- XLSX -----------------------------------------------------------------

def _xlsx_cells(director):
    wb = load_workbook(generate_quotation_xlsx(_doc(), CUSTOMER, [(_item(), None, [])], company=COMPANY, director=director))
    return [(c.coordinate, c.value) for row in wb.active.iter_rows() for c in row if c.value is not None]


def test_xlsx_supplier_block_shows_position():
    cells = _xlsx_cells(_director(position="MD"))
    labels = [v for _, v in cells]
    assert "Position" in labels
    coord = next(c for c, v in cells if v == "Position")
    assert ("B" + coord[1:], "MD") in cells


def test_xlsx_legacy_director_has_no_position_row():
    for d in (None, _director(), _director(position=None)):
        assert "Position" not in [v for _, v in _xlsx_cells(d)]


# ---- persistence / API / data flow ---------------------------------------

def test_position_lives_on_the_director_not_the_document():
    assert "position" in models.CompanyDirector.__table__.columns
    assert models.CompanyDirector.__table__.columns["position"].nullable
    assert "position" not in models.Document.__table__.columns
    for schema in (schemas.DocumentCreate, schemas.DocumentUpdate, schemas.DocumentOut):
        assert "position" not in schema.model_fields, f"{schema.__name__} must not duplicate position"


def test_director_schemas_carry_position_and_it_is_optional():
    assert schemas.CompanyDirectorCreate(name="X").position is None
    assert schemas.CompanyDirectorCreate(name="X", position="MD").position == "MD"
    assert schemas.CompanyDirectorUpdate(position="CEO").model_dump(exclude_unset=True) == {"position": "CEO"}
    out = schemas.CompanyDirectorOut.model_validate(SimpleNamespace(
        id=uuid.uuid4(), company_id=uuid.uuid4(), name="X", address=None, contact_no=None, email=None,
        position=None, seal_object_key=None, sort_order=1, created_at=__import__("datetime").datetime.now()))
    assert out.position is None


class _FakeQuery:
    def __init__(self, row): self.row = row
    def filter_by(self, **kw): return self
    def first(self): return self.row


class _FakeDB:
    def __init__(self, rows): self.rows = rows
    def query(self, model): return _FakeQuery(self.rows.get(model))


def _gather(director_row):
    doc = SimpleNamespace(id=uuid.uuid4(), customer_id=None, company_id=None, director_id=director_row and uuid.uuid4(),
                          items=[], project_id=None)
    db = _FakeDB({models.Document: doc, models.CompanyDirector: director_row})
    return documents_router._gather_export_data(doc.id, db)


def _director_row(position):
    return SimpleNamespace(name="Zaw Zaw", address=None, contact_no=None, email=None, position=position, seal_object_key=None)


def test_export_data_passes_the_selected_directors_position_through():
    director = _gather(_director_row("MD"))[8]
    assert director["position"] == "MD"
    assert _gather(_director_row(None))[8]["position"] is None
    assert _gather(None)[8] is None  # no MD selected on the document


def test_end_to_end_selected_director_to_pdf_text():
    director = _gather(_director_row("Operations Director"))[8]
    assert "Operations Director" in _supplier_block(_pdf_text(director))


# ---- migration ------------------------------------------------------------

def test_0004_migration_is_guarded_nullable_and_chained():
    path = pathlib.Path(__file__).resolve().parents[1] / "alembic" / "versions" / "0004_company_director_position.py"
    src = path.read_text()
    tree = ast.parse(src)
    up = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == "upgrade")
    body = ast.unparse(up)
    assert "not in _existing_columns('company_directors')" in body, "add_column must be guarded (idempotent)"
    assert "nullable=False" not in body and "server_default" not in body
    assert 'down_revision: Union[str, None] = "0003_document_deleted_by_at"' in src
    assert len("0004_director_position") <= 32  # alembic_version.version_num is VARCHAR(32)
