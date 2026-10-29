"""
Loop iteration 7 — a document can carry a one-time customer typed by hand
(documents.end_user_*) instead of a saved Customer (documents.customer_id).
The typed customer is stored on the document only — never as a Customer row
— and prints in the quotation's END USER block.

Real HTTP API + ORM on SQLite (JSONB mapped to JSON), same setup as
test_supplier_position_e2e.py; only search indexing and the audit log are
stubbed. The migration is run for real against SQLite, twice.
"""
import importlib.util
import io
import pathlib
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine, inspect, text
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import models
from app.core.config import settings
from app.core.db import Base, get_db
from app.main import app
from app.routers import documents as documents_router

MANUAL = {"end_user_name": "Daw Hla Hla", "end_user_organization": "Golden Lotus Trading",
          "end_user_address": "No. 5, Pansodan Road, Yangon", "end_user_phone": "09-111222333",
          "end_user_email": "hla@lotus.example"}


@compiles(JSONB, "sqlite")
def _jsonb_as_json(type_, compiler, **kw):
    return "JSON"


@pytest.fixture
def ctx(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False)

    def _get_db():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = _get_db
    monkeypatch.setattr(documents_router, "index_document", lambda doc: None)
    monkeypatch.setattr(documents_router, "log_action", lambda *a, **k: None)
    client = TestClient(app, headers={"X-Internal-Secret": settings.internal_shared_secret})
    yield client, Session
    app.dependency_overrides.clear()


def _create(client, **body):
    return client.post("/documents", json={"doc_type": "quotation", "issue_date": str(date(2026, 10, 29)),
                                           "items": [], **body})


def _edit(client, doc_id, **body):
    return client.patch(f"/documents/{doc_id}", json={"items": [], **body})


def _saved_customer(client):
    r = client.post("/customers", json={"name": "Saved Customer Co.",
                                        "billing_address": {"line1": "No. 9, Saved Street", "city": "Mandalay"}})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _end_user_pdf(client, doc_id):
    r = client.get(f"/documents/{doc_id}/export/quotation-pdf")
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    text_ = PdfReader(io.BytesIO(r.content)).pages[0].extract_text()
    return text_.split("END USER", 1)[1].split("No Item", 1)[0]


def _end_user_xlsx(client, doc_id):
    r = client.get(f"/documents/{doc_id}/export/quotation")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    return {ws.cell(r_, 5).value: ws.cell(r_, 6).value for r_ in range(1, 15) if ws.cell(r_, 5).value}


# ---- API / validation ------------------------------------------------------

def test_manual_customer_is_stored_on_the_document_and_never_saved_as_a_customer(ctx):
    client, Session = ctx
    before = Session().query(models.Customer).count()
    r = _create(client, **{k: f"  {v}  " for k, v in MANUAL.items()})  # whitespace is trimmed
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["customer_id"] is None
    assert {k: body[k] for k in MANUAL} == MANUAL
    assert Session().query(models.Customer).count() == before
    assert client.get(f"/documents/{body['id']}").json()["end_user_name"] == MANUAL["end_user_name"]
    assert any(d["end_user_name"] == MANUAL["end_user_name"] for d in client.get("/documents").json())


def test_saved_customer_and_manual_name_together_is_rejected(ctx):
    client, _ = ctx
    cid = _saved_customer(client)
    r = _create(client, customer_id=cid, end_user_name="Someone")
    assert r.status_code == 422 and "not both" in r.json()["detail"]
    doc = _create(client, customer_id=cid).json()
    assert _edit(client, doc["id"], customer_id=cid, end_user_name="Someone").status_code == 422


def test_blank_name_counts_as_not_sent_and_no_customer_stays_allowed(ctx):
    client, _ = ctx
    r = _create(client, end_user_name="   ", end_user_organization="Orphan Org")
    assert r.status_code == 201, r.text  # same as today: a document with no customer is allowed
    assert all(r.json()[k] is None for k in MANUAL)
    cid = _saved_customer(client)
    assert _create(client, customer_id=cid, end_user_name="  ").status_code == 201  # blank name + saved id is fine


def test_switching_between_manual_and_saved_customer(ctx):
    client, _ = ctx
    cid = _saved_customer(client)
    doc_id = _create(client, **MANUAL).json()["id"]

    to_saved = _edit(client, doc_id, customer_id=cid)
    assert to_saved.status_code == 200, to_saved.text
    assert to_saved.json()["customer_id"] == cid
    assert all(to_saved.json()[k] is None for k in MANUAL), "manual fields must be cleared"

    to_manual = _edit(client, doc_id, **MANUAL)
    assert to_manual.status_code == 200, to_manual.text
    assert to_manual.json()["customer_id"] is None
    assert to_manual.json()["end_user_name"] == MANUAL["end_user_name"]


# ---- exports ---------------------------------------------------------------

def test_exports_print_the_manual_customer(ctx):
    client, _ = ctx
    doc_id = _create(client, **MANUAL).json()["id"]
    pdf = _end_user_pdf(client, doc_id)
    assert "Daw Hla Hla" in pdf and "Organization Golden Lotus Trading" in pdf
    assert "Address No. 5, Pansodan Road, Yangon" in pdf
    assert "09-111222333" not in pdf and "hla@lotus.example" not in pdf  # stored only, not printed
    xlsx = _end_user_xlsx(client, doc_id)
    assert xlsx["Customer Name"] == "Daw Hla Hla"
    assert xlsx["Organization"] == "Golden Lotus Trading"
    assert xlsx["Address"] == "No. 5, Pansodan Road, Yangon"


def test_empty_organization_prints_the_name(ctx):
    client, _ = ctx
    doc_id = _create(client, end_user_name="Daw Hla Hla").json()["id"]
    assert "Organization Daw Hla Hla" in _end_user_pdf(client, doc_id)
    assert _end_user_xlsx(client, doc_id)["Organization"] == "Daw Hla Hla"


def test_saved_customer_end_user_output_is_unchanged(ctx):
    client, _ = ctx
    doc_id = _create(client, customer_id=_saved_customer(client)).json()["id"]
    pdf = _end_user_pdf(client, doc_id)
    assert "Saved Customer Co." in pdf
    assert "Organization Saved Customer Co." in pdf  # still the customer name, as before
    assert "Address No. 9, Saved Street, Mandalay" in pdf
    xlsx = _end_user_xlsx(client, doc_id)
    xlsx.pop("Quotation ID")  # auto-generated doc number
    assert xlsx == {"END USER": None, "Customer Name": "Saved Customer Co.", "Organization": "Saved Customer Co.",
                    "Address": "No. 9, Saved Street, Mandalay", "Price (USD)": "Amount (USD)"}


# ---- migration -------------------------------------------------------------

def _load_0005():
    path = pathlib.Path(__file__).resolve().parents[1] / "alembic" / "versions" / "0005_document_manual_end_user.py"
    spec = importlib.util.spec_from_file_location("m0005", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_0005_migration_runs_twice_safely_and_keeps_rows():
    from alembic.migration import MigrationContext
    from alembic.operations import Operations

    m0005 = _load_0005()
    assert m0005.down_revision == "0004_director_position" and len(m0005.revision) <= 32
    engine = create_engine("sqlite://")
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE documents (id VARCHAR PRIMARY KEY, doc_number VARCHAR)"))
        conn.execute(text("INSERT INTO documents VALUES ('d1', 'QT-OLD')"))
        with Operations.context(MigrationContext.configure(conn)):
            m0005.upgrade()
            m0005.upgrade()  # second run: every column already exists -> no-op, no error
        cols = [c["name"] for c in inspect(conn).get_columns("documents")]
        assert cols == ["id", "doc_number", "end_user_name", "end_user_organization",
                        "end_user_address", "end_user_phone", "end_user_email"]
        assert conn.execute(text("SELECT doc_number, end_user_name FROM documents")).all() == [("QT-OLD", None)]
