"""
Loop iteration 6 — a bold "Regards," line directly above the MD's seal (or
above the name/position fallback when no seal is uploaded), horizontally
centered on it, in both exports. Only shown when a director is selected.

PDF geometry is read from WeasyPrint's box tree via the `layout` fixture
imported from test_quotation_visual_style.py; the page sweep reuses
test_quotation_pdf_layout.py's generator and its MD-orphan check.
"""
import pytest
from openpyxl import load_workbook
from openpyxl.utils.units import EMU_to_pixels

from app.core.export_quotation import generate_quotation_xlsx
from tests.test_quotation_pdf_layout import _assert_md_page_not_orphaned, _generate
from tests.test_quotation_visual_style import (  # noqa: F401  (layout is a fixture)
    COMPANY,
    CUSTOMER,
    DIRECTOR,
    MM,
    _document,
    _images,
    _line_item,
    _png,
    layout,
)

MD_SEAL_PX = 130  # .md-seal; the header logo/seal are 120px


def _regards(page):
    return [b for b in page if getattr(b, "text", None) == "Regards,"]


def _center_x(box):
    return box.position_x + box.width / 2


def _md_seal(page):
    return [i for i in _images(page) if i.width == pytest.approx(MD_SEAL_PX, abs=0.5)]


def test_pdf_regards_is_bold_above_the_seal_and_centered_on_it(layout):
    result = layout(n_items=3, director_seal_bytes=_png())
    found = [(n, b) for n, page in enumerate(result.pages) for b in _regards(page)]
    assert len(found) == 1, f"'Regards,' should appear exactly once, found {len(found)}"
    n, text = found[0]
    seals = _md_seal(result.pages[n])
    assert len(seals) == 1, "'Regards,' must be on the same page as the MD seal"
    seal = seals[0]
    assert text.style["font_weight"] >= 700, "'Regards,' must be bold"
    assert text.position_y + text.height <= seal.position_y, "'Regards,' must sit above the seal"
    assert _center_x(text) == pytest.approx(_center_x(seal), abs=1 * MM), (
        f"'Regards,' center is {(_center_x(text) - _center_x(seal)) / MM:+.2f}mm off the seal's center"
    )


def test_pdf_without_a_seal_regards_is_centered_above_the_md_name(layout):
    result = layout(n_items=3)  # director, no seal -> name/position fallback
    page = next(p for p in result.pages if _regards(p))
    text = _regards(page)[0]
    name = [b for b in page if getattr(b, "text", None) == DIRECTOR["name"] and b.position_y > text.position_y]
    assert name, "the MD name must render below 'Regards,'"
    assert text.position_y + text.height <= name[0].position_y
    assert _center_x(text) == pytest.approx(_center_x(name[0]), abs=1 * MM)


def test_pdf_without_a_director_has_no_regards(layout):
    result = layout(n_items=3, director=None)
    assert not any(_regards(p) for p in result.pages)


@pytest.mark.parametrize("n_items", [1, 3, 5, 8, 12, 20, 30, 80])
def test_regards_and_seal_always_share_a_page(layout, n_items):
    result = layout(n_items=n_items, director_seal_bytes=_png())
    regards_pages = [n for n, p in enumerate(result.pages) if _regards(p)]
    seal_pages = [n for n, p in enumerate(result.pages) if _md_seal(p)]
    assert regards_pages == seal_pages and len(seal_pages) == 1, (
        f"with {n_items} items: 'Regards,' on page(s) {regards_pages}, seal on page(s) {seal_pages}"
    )


@pytest.mark.parametrize("n_items", [1, 3, 5, 8, 12, 20, 30, 80])
def test_md_block_with_regards_is_still_never_orphaned(n_items):
    for text in _generate(n_items):
        if "Managing Director" in text:
            _assert_md_page_not_orphaned(text, n_items)


# ---- XLSX -----------------------------------------------------------------

def _xlsx(**kw):
    return load_workbook(generate_quotation_xlsx(_document(), CUSTOMER, [(_line_item(1), None, [])],
                                                 company=COMPANY, **kw)).active


def _regards_cells(ws):
    return [c for row in ws.iter_rows() for c in row if c.value == "Regards,"]


def test_xlsx_regards_is_bold_above_the_seal_in_the_seals_column_and_centered():
    ws = _xlsx(director=DIRECTOR, director_seal_bytes=_png())
    cells = _regards_cells(ws)
    assert len(cells) == 1
    cell = cells[0]
    assert cell.font.b and cell.alignment.horizontal == "center"
    seal = next(i for i in ws._images if i.anchor._from.col == 6)  # column G
    assert cell.column_letter == "G", "'Regards,' must be in the seal's column"
    assert cell.row == seal.anchor._from.row, "'Regards,' must be on the row directly above the seal (anchor rows are 0-based)"
    # The seal is centered in G (same 210px width Excel shows for width 30),
    # so its center is the centered text's center.
    col_px = int((256 * ws.column_dimensions["G"].width + 128 // 7) / 256 * 7)
    seal_center = EMU_to_pixels(seal.anchor._from.colOff) + EMU_to_pixels(seal.anchor.ext.width) / 2
    assert seal_center == pytest.approx(col_px / 2, abs=1)


def test_xlsx_without_a_seal_regards_is_centered_over_the_name():
    ws = _xlsx(director=DIRECTOR)
    cell = _regards_cells(ws)[0]
    assert cell.font.b and cell.alignment.horizontal == "center"
    assert f"F{cell.row}:G{cell.row}" in {str(r) for r in ws.merged_cells.ranges}
    assert ws.cell(cell.row + 1, 6).value == DIRECTOR["name"]


def test_xlsx_without_a_director_has_no_regards():
    assert not _regards_cells(_xlsx(director=None, director_seal_bytes=_png()))
