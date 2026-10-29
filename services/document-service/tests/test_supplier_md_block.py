"""
Loop iteration 1 — the quotation's SUPPLIER block names the company on its
heading line ("SUPPLIER: <company>") and describes the Managing Director
selected on the document (Document.director_id) in the rows below it:

    SUPPLIER: Swift Solutions Co., Ltd.
    Name        U Wathone Htoo        <- director.name, "—" if no MD
    Position    Managing Director     <- director.position, else company.position, else row omitted
    Address     BUILDING NO.4, ...    <- director.address, else company.address
    Contact No  09-988733966          <- director.contact_no, else company.contact_no

There is no Email row in the block any more; the company's support email
stays in the Customer Service footer. Covered for both exports (PDF text via
pypdf, XLSX cells via openpyxl), calling the generators directly with
lightweight fakes — no Postgres/Docker needed.
"""
import uuid
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.export_pdf import generate_quotation_pdf
from app.core.export_quotation import generate_quotation_xlsx

COMPANY = {"name": "Swift Solutions Co., Ltd.", "position": None, "address": "No. 7 Company Road, Yangon",
           "contact_no": "01-555-0100", "support_email": "support@swift.example"}
CUSTOMER = {"id": str(uuid.uuid4()), "name": "Beta Customer", "billing_address": {"line1": "Street 9"}}
FULL_MD = {"name": "U Wathone Htoo", "position": "Managing Director", "address": "BUILDING NO.4, MD Street",
           "contact_no": "09-988733966", "email": "md@swift.example"}
BARE_MD = {"name": "U Wathone Htoo", "position": "Managing Director", "address": None,
           "contact_no": "", "email": None}


def _doc():
    return SimpleNamespace(issue_date=date(2026, 1, 1), currency="USD", doc_number="QT-1",
                           terms_and_conditions=None, subtotal=None, tax_total=None, total=None)


def _item():
    return SimpleNamespace(description="Widget", quantity=1, unit_price=10, unit="Nos", remark="",
                           product_id=None, tax_rate=None)


# ---- PDF ------------------------------------------------------------------

def _pdf_supplier(director, company=COMPANY):
    """(heading line, {label: value}) of the PDF's Supplier block."""
    pdf = generate_quotation_pdf(_doc(), CUSTOMER, [(_item(), None, [])], company=company, director=director)
    text = PdfReader(BytesIO(pdf.getvalue())).pages[0].extract_text()
    block = "SUPPLIER" + text.split("SUPPLIER", 1)[1].split("END USER", 1)[0]
    heading, *lines = block.strip().split("\n")
    rows = {}
    for line in lines:
        for label in ("Name", "Position", "Address", "Contact No", "Email"):
            if line.startswith(label + " ") or line == label:
                rows[label] = line[len(label):].strip()
    return heading.strip(), rows, block


def test_pdf_director_with_all_fields():
    heading, rows, block = _pdf_supplier(FULL_MD)
    assert heading == "SUPPLIER: Swift Solutions Co., Ltd."
    assert rows == {
        "Name": "U Wathone Htoo",
        "Position": "Managing Director",
        "Address": "BUILDING NO.4, MD Street",
        "Contact No": "09-988733966",
    }
    assert "Email" not in block and "support@swift.example" not in block
    # the company's own address/contact are not shown when the MD has their own
    assert "Company Road" not in block and "01-555-0100" not in block


def test_pdf_director_without_address_or_contact_falls_back_to_the_company():
    heading, rows, _ = _pdf_supplier(BARE_MD)
    assert heading == "SUPPLIER: Swift Solutions Co., Ltd."
    assert rows["Name"] == "U Wathone Htoo"
    assert rows["Address"] == "No. 7 Company Road, Yangon"
    assert rows["Contact No"] == "01-555-0100"


def test_pdf_no_director_shows_a_dash_and_the_company_details():
    heading, rows, block = _pdf_supplier(None)
    assert heading == "SUPPLIER: Swift Solutions Co., Ltd."
    assert rows["Name"] == "—"
    assert "Position" not in rows  # neither MD nor company position -> row omitted
    assert rows["Address"] == "No. 7 Company Road, Yangon"
    assert rows["Contact No"] == "01-555-0100"
    assert "Email" not in block


def test_pdf_no_company_keeps_the_bare_heading():
    heading, rows, _ = _pdf_supplier(None, company=None)
    assert heading == "SUPPLIER"
    assert rows["Name"] == "—"


def test_pdf_footer_still_has_the_company_support_email():
    pdf = generate_quotation_pdf(_doc(), CUSTOMER, [(_item(), None, [])], company=COMPANY, director=FULL_MD)
    text = "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf.getvalue())).pages)
    assert "Email: support@swift.example" in text


# ---- XLSX -----------------------------------------------------------------

def _xlsx_supplier(director, company=COMPANY):
    """(heading cell A, {label: value}) of the XLSX Supplier block — the A/B
    cells from the SUPPLIER heading row down to the first empty A cell."""
    ws = load_workbook(generate_quotation_xlsx(_doc(), CUSTOMER, [(_item(), None, [])],
                                               company=company, director=director)).active
    start = next(r for r in range(1, ws.max_row + 1)
                 if isinstance(ws.cell(r, 1).value, str) and ws.cell(r, 1).value.startswith("SUPPLIER"))
    assert ws.cell(start, 2).value is None, "the heading is one string in A; B must stay empty"
    rows = {}
    r = start + 1
    while ws.cell(r, 1).value:
        rows[ws.cell(r, 1).value] = ws.cell(r, 2).value
        r += 1
    return ws.cell(start, 1).value, rows


def test_xlsx_director_with_all_fields():
    heading, rows = _xlsx_supplier(FULL_MD)
    assert heading == "SUPPLIER: Swift Solutions Co., Ltd."
    assert rows == {
        "Name": "U Wathone Htoo",
        "Position": "Managing Director",
        "Address": "BUILDING NO.4, MD Street",
        "Contact No": "09-988733966",
    }  # no Email row


def test_xlsx_director_without_address_or_contact_falls_back_to_the_company():
    _, rows = _xlsx_supplier(BARE_MD)
    assert rows["Name"] == "U Wathone Htoo"
    assert rows["Address"] == "No. 7 Company Road, Yangon"
    assert rows["Contact No"] == "01-555-0100"


def test_xlsx_no_director_shows_a_dash_and_the_company_details():
    heading, rows = _xlsx_supplier(None)
    assert heading == "SUPPLIER: Swift Solutions Co., Ltd."
    assert rows == {"Name": "—", "Address": "No. 7 Company Road, Yangon", "Contact No": "01-555-0100"}


def test_xlsx_no_company_keeps_the_bare_heading():
    heading, rows = _xlsx_supplier(None, company=None)
    assert heading == "SUPPLIER"
    assert rows["Name"] == "—"


def test_xlsx_director_values_are_still_formula_escaped():
    _, rows = _xlsx_supplier(dict(FULL_MD, name="=HYPERLINK(\"x\")", contact_no="+95 9 123"))
    assert rows["Name"].startswith("'=")
    assert rows["Contact No"] == "'+95 9 123"
