"""
Regression test for the same tax-display bug fixed in the PDF export
(see test_quotation_pdf_layout.py and docs/SECURITY_HARDENING_LOG.md):
the XLSX quotation export's "Total:" row was a plain formula summing
each line's Amount column, with no trace of tax anywhere in the sheet.

generate_quotation_xlsx now adds Subtotal/Tax/Total rows (as live
formulas, so editing a line item's qty/price in Excel still recalculates
correctly) whenever the document has a tax rate, and falls back to a
single Total row (unchanged from before) when it doesn't.
"""
from datetime import date
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.core.export_quotation import generate_quotation_xlsx

CUSTOMER = {
    "id": "11111111-1111-1111-1111-111111111111",
    "name": "Beta Customer Co., Ltd.",
    "billing_address": {"line1": "No. 99, Customer Street", "city": "Mandalay"},
}


def _document():
    return SimpleNamespace(
        issue_date=date(2026, 1, 1),
        currency="USD",
        doc_number="QT-0001",
        terms_and_conditions=None,
    )


def _line_item(name, qty=1, price=100, tax_rate=None):
    return SimpleNamespace(
        description=name,
        quantity=qty,
        unit_price=price,
        unit="Nos",
        remark="",
        product_id=None,
        tax_rate=tax_rate,
    )


def _labels_and_formulas(ws):
    """Returns {label: (row, formula_or_value)} for every row whose
    column A holds a recognizable totals-section label."""
    found = {}
    for row in ws.iter_rows():
        label_cell = row[0]
        if label_cell.value in ("Subtotal:", "Total:") or (
            isinstance(label_cell.value, str) and label_cell.value.startswith("Tax (")
        ):
            found[label_cell.value] = (label_cell.row, ws.cell(row=label_cell.row, column=6).value)
    return found


def test_xlsx_shows_tax_breakdown_when_items_have_tax():
    items = [(_line_item(f"Item {i}", tax_rate=10), None, []) for i in range(1, 4)]
    wb = generate_quotation_xlsx(_document(), CUSTOMER, items)
    ws = load_workbook(filename=BytesIO(wb.getvalue())).active

    found = _labels_and_formulas(ws)
    assert "Subtotal:" in found, "expected a Subtotal row when the document has tax"
    tax_label = next((k for k in found if k.startswith("Tax (")), None)
    assert tax_label == "Tax (10%):", f"expected a 'Tax (10%):' row, found labels: {list(found)}"
    assert "Total:" in found

    subtotal_row, subtotal_formula = found["Subtotal:"]
    tax_row, tax_formula = found[tax_label]
    total_row, total_formula = found["Total:"]

    assert isinstance(subtotal_formula, str) and subtotal_formula.startswith("=")
    assert tax_formula == f"=F{subtotal_row}*10.0/100"
    assert total_formula == f"=F{subtotal_row}+F{tax_row}"
    assert tax_row == subtotal_row + 1
    assert total_row == tax_row + 1


def test_xlsx_omits_tax_breakdown_when_items_have_no_tax():
    items = [(_line_item(f"Item {i}", tax_rate=0), None, []) for i in range(1, 4)]
    wb = generate_quotation_xlsx(_document(), CUSTOMER, items)
    ws = load_workbook(filename=BytesIO(wb.getvalue())).active

    found = _labels_and_formulas(ws)
    assert "Subtotal:" not in found
    assert not any(k.startswith("Tax (") for k in found)
    assert "Total:" in found
